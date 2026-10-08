"""O sistema lê a estrutura que a extração de texto joga fora. A IA julga o que isso significa; aqui só se prova
que os fatos objetivos chegam ao modelo e que a falta de dado nunca é omitida."""

import io
import zipfile

import docx
import pytest
from PIL import Image
from pptx import Presentation

from backend.src.services.chat_tools import _read_docx_text, _read_epub_text, _read_pdf_text, _read_pptx_text
from backend.src.services.document_facts import docx_facts, epub_facts, pdf_facts, pptx_facts


def _png() -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (20, 20), (200, 30, 30)).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def _docx_bytes(with_alt: bool = False) -> bytes:
    document = docx.Document()
    document.add_paragraph("Relatório")
    document.add_picture(_png())
    if with_alt:
        document.element.body.xpath(".//wp:docPr")[0].set("descr", "Gráfico de vendas em alta")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Mês"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_reports_missing_title_missing_alt_and_unmarked_table_header():
    text = docx_facts(_docx_bytes())
    assert "Título nas propriedades do documento: nenhum" in text
    assert "sem texto alternativo: 1" in text
    assert "primeira linha marcada como cabeçalho: não" in text


def test_docx_reports_the_alt_text_that_exists():
    text = docx_facts(_docx_bytes(with_alt=True))
    assert "com texto alternativo: 1" in text
    assert "Gráfico de vendas em alta" in text


def test_docx_reader_prepends_the_facts_to_the_extracted_text():
    text = _read_docx_text(io.BytesIO(_docx_bytes()))
    assert text.startswith("[Fatos observados na estrutura do arquivo")
    assert "[Style: Normal] Relatório" in text


def test_unreadable_docx_says_so_instead_of_staying_silent():
    assert "Não foi possível ler a estrutura do DOCX" in docx_facts(b"nao e um docx")


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(_png(), 0, 0)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_pptx_reports_missing_slide_title_and_shows_the_alt_text_even_when_it_is_only_a_file_name():
    text = pptx_facts(_pptx_bytes())
    assert "sem título de slide" in text
    assert "texto alternativo = " in text  # o valor aparece, para a IA julgar se é só nome de arquivo
    assert _read_pptx_text(io.BytesIO(_pptx_bytes())).startswith("[Fatos observados")


def _pdf_bytes(extra_catalog: str = "") -> bytes:
    body = "BT /F1 12 Tf 72 700 Td (Ola mundo) Tj ET"
    objs = [
        f"<< /Type /Catalog /Pages 2 0 R {extra_catalog} >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(body)} >>\nstream\n{body}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objs, 1):
        offsets.append(len(data))
        data += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(data)
    data += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        data += f"{offset:010d} 00000 n \n".encode()
    data += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return data


def test_pdf_reports_untagged_without_language_or_title():
    text = pdf_facts(_pdf_bytes())
    assert "Marcado (tagged PDF, /MarkInfo /Marked): não" in text
    assert "Idioma do documento (/Lang): nenhum" in text
    assert "Título nos metadados: nenhum" in text
    assert _read_pdf_text(io.BytesIO(_pdf_bytes())).startswith("[Fatos observados")


def test_pdf_reports_tagged_with_language():
    text = pdf_facts(_pdf_bytes("/MarkInfo << /Marked true >> /Lang (pt-BR)"))
    assert "Marcado (tagged PDF, /MarkInfo /Marked): sim" in text
    assert "'pt-BR'" in text


def _epub_bytes() -> bytes:
    opf = (
        '<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>Livro</dc:title></metadata><manifest>"
        '<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
        '<spine><itemref idref="c1"/></spine></package>'
    )
    chapter = "<html><body><h1>Capítulo</h1><img src='a.png'/><table><tr><td>x</td></tr></table></body></html>"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/c1.xhtml", chapter)
    return buffer.getvalue()


def test_epub_reports_missing_language_navigation_alt_and_table_headers():
    text = epub_facts(_epub_bytes())
    assert "Idioma (dc:language): 'nenhum'" in text
    assert "Documento de navegação (nav): ausente" in text
    assert "sem atributo alt: 1" in text
    assert "células de cabeçalho th: 0" in text
    assert "schema:accessibilitySummary: ausente" in text
    assert _read_epub_text(io.BytesIO(_epub_bytes())).startswith("[Fatos observados")


@pytest.mark.parametrize("reader", [docx_facts, pptx_facts, pdf_facts, epub_facts])
def test_every_reader_reports_what_it_could_not_read(reader):
    assert "Não foi possível ler" in reader(b"lixo") or "não tem o manifesto" in reader(b"lixo")


def test_pdf_reader_does_not_hang_on_a_mocked_or_odd_reader():
    """Regressão (a suíte inteira estourou a memória): um leitor trocado por objeto de mentira fazia o percurso da
    árvore de estrutura não terminar. O percurso só segue dicionários de PDF de verdade, cada um uma vez."""
    from unittest.mock import MagicMock, patch

    with patch("pypdf.PdfReader", return_value=MagicMock()):
        text = pdf_facts(b"%PDF qualquer coisa")
    assert text.startswith("[Fatos observados")


def test_pdf_structure_walk_counts_figures_with_and_without_alt():
    import pypdf
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject

    def figure(alt: str | None) -> DictionaryObject:
        node = DictionaryObject({NameObject("/S"): NameObject("/Figure")})
        if alt is not None:
            node[NameObject("/Alt")] = TextStringObject(alt)
        return node

    cyclic = DictionaryObject({NameObject("/S"): NameObject("/Document")})
    cyclic[NameObject("/K")] = ArrayObject([figure("Gráfico"), figure(None), cyclic])  # inclui um ciclo
    fake_root = {"/StructTreeRoot": cyclic, "/Lang": TextStringObject("pt-BR")}

    reader = pypdf.PdfReader(io.BytesIO(_pdf_bytes()))
    reader.trailer["/Root"].update({NameObject(k): v for k, v in fake_root.items()})
    from unittest.mock import patch

    with patch("pypdf.PdfReader", return_value=reader):
        text = pdf_facts(b"x")
    assert "Figuras na estrutura: 2; com texto alternativo (/Alt): 1; sem: 1" in text
