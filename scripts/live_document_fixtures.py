"""Arquivos de teste com defeitos CONHECIDOS para o módulo 4 (documentos e imagens), usados por verify_live_modules_1_to_3.py.

Cada construtor devolve os bytes do arquivo; `EXPECTED` traz a verdade conhecida (id -> defeito) que o juiz de IA usa para dizer
o que a resposta trouxe e o que faltou. O arquivo é gerado na hora: nada de binário guardado no repositório.
"""

from __future__ import annotations

import base64
import io
import zipfile

import docx
import openpyxl
from PIL import Image, ImageDraw
from pptx import Presentation


def _png(color: tuple[int, int, int] = (200, 30, 30)) -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 40), color).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def docx_bytes() -> bytes:
    document = docx.Document()
    document.core_properties.title = ""
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Relatório de vendas do trimestre")
    run.bold = True
    run.font.size = docx.shared.Pt(20)  # parece título, mas não usa o estilo de título
    document.add_paragraph("As vendas cresceram em todas as regiões.")
    document.add_picture(_png())
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Mês"
    table.cell(0, 1).text = "Total"
    table.cell(1, 0).text = "Janeiro"
    table.cell(1, 1).text = "10"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(_png(), 0, 0)
    box = slide.shapes.add_textbox(0, 0, 3000000, 500000)
    box.text_frame.text = "Resultados do trimestre"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Vendas por região"
    sheet.append(["Norte", 10, 12])  # os dados começam sem linha de cabeçalho
    sheet.append(["Sul", 8, 9])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pdf_bytes() -> bytes:
    body = "BT /F1 12 Tf 72 700 Td (Relatorio de vendas) Tj ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(body)} >>\nstream\n{body}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    data = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objs, 1):
        offsets.append(len(data))
        data += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(data)
    data += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        data += f"{offset:010d} 00000 n \n".encode()
    data += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return data


def epub_bytes() -> bytes:
    opf = (
        '<package xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>Livro</dc:title></metadata><manifest>"
        '<item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
        '<spine><itemref idref="c1"/></spine></package>'
    )
    chapter = "<html><body><h1>Capítulo</h1><img src='a.png'/><table><tr><td>x</td></tr></table></body></html>"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/c1.xhtml", chapter)
    return buffer.getvalue()


def banner_png() -> bytes:
    image = Image.new("RGB", (640, 200), (255, 255, 255))
    ImageDraw.Draw(image).text((40, 80), "PROMOCAO 50% DE DESCONTO", fill=(204, 204, 204))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def corrupted_docx() -> bytes:
    return b"isto nao e um arquivo docx de verdade" * 20


def attach(name: str, data: bytes) -> str:
    return f"=== {name} ===\n{base64.b64encode(data).decode()}\n"


EXPECTED: dict[str, dict[str, str]] = {
    "docx": {
        "title": "o documento não tem título nas propriedades",
        "language": "o documento não declara o idioma",
        "img-alt": "a imagem não tem texto alternativo",
        "fake-heading": "'Relatório de vendas do trimestre' parece título (negrito, tamanho grande) mas não usa o estilo de título",
        "table-header": "a tabela não marca a primeira linha como cabeçalho",
    },
    "pptx": {
        "slide-title": "o slide não tem título",
        "img-alt": "a imagem não tem texto alternativo útil (só o nome do arquivo, se houver)",
        "reading-order": "a ordem de leitura dos objetos do slide não foi garantida (imagem antes do texto)",
    },
    "xlsx": {
        "merged-cells": "há células mescladas (A1:C1)",
        "default-sheet-name": "a planilha mantém o nome padrão 'Sheet'",
        "no-header-row": "os dados começam sem uma linha de cabeçalho de colunas",
    },
    "pdf": {
        "untagged": "o PDF não é marcado (tagged), então não tem estrutura para leitor de tela",
        "language": "o PDF não declara o idioma",
        "title": "o PDF não tem título nos metadados",
    },
    "epub": {
        "language": "o EPUB não declara o idioma",
        "navigation": "o EPUB não tem documento de navegação",
        "img-alt": "a imagem do capítulo não tem atributo alt",
        "table-header": "a tabela não tem células de cabeçalho (th)",
    },
    "imagem": {
        "text": "a imagem contém o texto 'PROMOCAO 50% DE DESCONTO'",
        "low-contrast": "o texto é cinza claro sobre fundo branco, com contraste muito baixo",
    },
}
